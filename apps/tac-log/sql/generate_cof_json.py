import openpyxl, re, json, sys

SRC = sys.argv[1] if len(sys.argv) > 1 else "Courses of Fire.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "courses-of-fire.patch.json"

def clean_str(v):
    if v is None:
        return None
    return str(v).strip()

wb = openpyxl.load_workbook(SRC, data_only=True)

courses = []

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    rows = list(ws.iter_rows(values_only=True))

    name = clean_str(rows[0][0])
    code = None
    total_rounds = None
    target_type = None
    for r in rows[:8]:
        if r[0] == 'Document ID:':
            code = clean_str(r[2])
        elif r[0] == 'Total Rounds:':
            total_rounds = r[2]
        elif r[0] == 'Target Type:':
            target_type = clean_str(r[2])

    phases = []
    zones = []
    in_scoring = False
    zone_table_started = False

    i = 0
    n = len(rows)
    while i < n:
        r = rows[i]
        c0 = r[0]
        if isinstance(c0, str) and c0.startswith("Phase "):
            m = re.match(r"Phase (\d+):\s*(.*)", c0)
            pnum = int(m.group(1)) if m else len(phases) + 1
            phases.append({"phase_number": pnum, "title": clean_str(c0), "phase_total_rounds": None, "strings": []})
            i += 1
            continue
        if c0 == 'String':
            i += 1
            continue
        if c0 in ('PHASE TOTAL:', 'PHASE TOTALS:'):
            if phases:
                total_val = r[2]
                phases[-1]["phase_total_rounds"] = total_val if isinstance(total_val, (int, float)) else None
            i += 1
            continue
        if c0 == 'COF TOTALS:':
            i += 1
            continue
        if c0 in ('Scoring Card', 'Scoring Section'):
            in_scoring = True
            zone_table_started = False
            i += 1
            continue
        if in_scoring and c0 == 'Zone':
            zone_table_started = True
            i += 1
            continue
        if in_scoring and zone_table_started:
            if c0 is None or c0 == 'Grader Name (Last, First, MI)':
                if c0 == 'Grader Name (Last, First, MI)':
                    in_scoring = False
                i += 1
                continue
            label = clean_str(c0)
            value = r[2]
            if label is not None and isinstance(value, (int, float)):
                zones.append({"zone_label": label, "value": value})
            i += 1
            continue
        if in_scoring:
            i += 1
            continue

        c1 = r[1]
        is_option_marker = (
            (isinstance(c0, str) and c0.strip().upper().startswith("OPTION")) or
            (c0 is None and isinstance(c1, str) and c1.strip().upper().startswith("OPTION"))
        )
        if is_option_marker and phases:
            option_label = clean_str(c0 if isinstance(c0, str) else c1)
            i += 1
            while i < n and isinstance(rows[i][0], (int, float)):
                rr = rows[i]
                has_weapon_col = clean_str(rr[2]) in ("HANDGUN", "RIFLE", "HG/RFL", "SHOTGUN") if isinstance(rr[2], str) else False
                if has_weapon_col:
                    distance, weapon, rounds_, time_limit, position, action = rr[1], rr[2], rr[3], rr[4], rr[5], rr[6]
                else:
                    distance, weapon, rounds_, time_limit, position, action = rr[1], None, rr[2], rr[3], rr[4], rr[5]
                phases[-1]["strings"].append({
                    "string_number": rr[0],
                    "option_label": option_label,
                    "distance": clean_str(distance),
                    "weapon": clean_str(weapon),
                    "rounds": clean_str(rounds_) if not isinstance(rounds_, (int, float)) else rounds_,
                    "time_limit": clean_str(time_limit),
                    "position": clean_str(position),
                    "action": clean_str(action),
                })
                i += 1
                if i < n and isinstance(rows[i][0], str) and rows[i][0].strip().upper().startswith("OPTION"):
                    break
            continue

        if isinstance(c0, (int, float)) and phases:
            has_weapon_col = clean_str(r[2]) in ("HANDGUN", "RIFLE", "HG/RFL", "SHOTGUN") if isinstance(r[2], str) else False
            if has_weapon_col:
                distance, weapon, rounds_, time_limit, position, action = r[1], r[2], r[3], r[4], r[5], r[6]
            else:
                distance, weapon, rounds_, time_limit, position, action = r[1], None, r[2], r[3], r[4], r[5]
            phases[-1]["strings"].append({
                "string_number": c0,
                "option_label": None,
                "distance": clean_str(distance),
                "weapon": clean_str(weapon),
                "rounds": clean_str(rounds_) if not isinstance(rounds_, (int, float)) else rounds_,
                "time_limit": clean_str(time_limit),
                "position": clean_str(position),
                "action": clean_str(action),
            })
            i += 1
            continue
        i += 1

    courses.append({
        "code": code,
        "name": name,
        "total_rounds": total_rounds,
        "target_type": target_type,
        "phases": phases,
        "scoring_zones": zones,
    })

with open(OUT, "w") as f:
    json.dump({"courses": courses}, f, indent=2)

print(f"wrote {OUT} with {len(courses)} courses")
